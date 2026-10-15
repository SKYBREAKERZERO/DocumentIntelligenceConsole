"""Native DOCX/PDF and scoped atomic source contract tests; synthetic evidence."""
from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import fitz
import pytest
from docx import Document

from app.engineering_intelligence.atomic import extract_atomic_candidates
from app.engineering_intelligence.condition_scope import (
    extract_requirement_scope, validate_requirement_scope,
)
from app.engineering_intelligence.native_layout import (
    inspect_docx_native, inspect_pdf_native, inspect_native_file,
)
from app.engineering_intelligence.requirements import extract_requirement_units


def _source(text: str):
    return extract_requirement_units(document_id="DOC-1", content_id=10,
                                     section_id="4.1.2", text=text)[0]


def test_explicit_english_condition_produces_exact_local_source_spans():
    unit = _source("When Bluetooth is connected, ECU shall notify Body ECU.")
    atom, = extract_atomic_candidates(unit)
    frame = extract_requirement_scope(unit, atom)
    assert frame.status == "EXPLICIT_PREFIX_CANDIDATE"
    assert frame.condition.exact_text == "When Bluetooth is connected,"
    assert frame.action.exact_text == " ECU shall notify Body ECU."
    assert frame.modality == "MANDATORY"
    validate_requirement_scope(unit, atom, frame)


def test_chinese_conditional_scope_keeps_full_width_comma():
    unit = _source("如果设备连接成功，系统应当通知Body ECU。")
    atom, = extract_atomic_candidates(unit)
    frame = extract_requirement_scope(unit, atom)
    assert frame.condition.exact_text == "如果设备连接成功，"
    assert frame.action.exact_text == "系统应当通知Body ECU。"


def test_unless_is_exception_not_both_conditions():
    unit = _source("Unless the region is NA, HU shall save the user profile.")
    atom, = extract_atomic_candidates(unit)
    frame = extract_requirement_scope(unit, atom)
    assert frame.exception and frame.exception.exact_text.startswith("Unless")
    assert frame.condition is None
    assert frame.action and frame.status == "EXPLICIT_PREFIX_CANDIDATE"


def test_undelimited_conditional_is_not_assigned_to_action():
    unit = _source("When BT connects ECU shall load profile")
    atom, = extract_atomic_candidates(unit)
    frame = extract_requirement_scope(unit, atom)
    assert frame.status == "UNRESOLVED"
    assert frame.condition is None and frame.action is None
    assert "SCOPE_BOUNDARY_NOT_EXPLICIT" in frame.review_reasons


def test_multi_atomic_scope_does_not_inherit_condition():
    unit = _source("When BT connects, HU shall load profile; HU must notify ECU.")
    a, b = extract_atomic_candidates(unit)
    assert extract_requirement_scope(unit, a).status == "UNRESOLVED"
    assert "PARENT_ATOMIC_SCOPE_UNRESOLVED" in extract_requirement_scope(unit, b).review_reasons


def test_scope_rejects_source_drift():
    unit = _source("ECU shall notify.")
    atom, = extract_atomic_candidates(unit)
    result = extract_requirement_scope(unit, atom)
    with pytest.raises(ValueError, match="does not match"):
        validate_requirement_scope(unit, atom, replace(result, action=replace(result.action, exact_text="fake")))


def test_docx_native_physical_grid_merges_and_no_fake_page(tmp_path: Path):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("BT specification")
    tab = doc.add_table(rows=3, cols=3)
    tab.cell(0, 0).text = "Region"
    tab.cell(0, 1).text = "Operations"
    tab.cell(0, 1).merge(tab.cell(0, 2))
    tab.cell(1, 0).text = "NA"
    tab.cell(1, 0).merge(tab.cell(2, 0))
    tab.cell(1, 1).text = "shall load"
    doc.save(path)
    out = inspect_docx_native(path)
    assert out.format == "DOCX"
    assert len(out.tables) == 1
    table, = out.tables
    assert table.page_number is None
    assert table.grid_columns == 3
    assert table.cells[1].column_span == 2
    master = next(c for c in table.cells if c.row == 1 and c.column == 0)
    assert master.row_span == 2
    cont = next(c for c in table.cells if c.row == 2 and c.column == 0)
    assert cont.is_merge_continuation and cont.merge_origin == (1, 0)
    assert "HORIZONTAL_MERGE_PRESENT" in table.review_reasons
    assert "TABLE_HEADER_NOT_AUTOMATICALLY_VERIFIED" in table.review_reasons
    assert "DOCX_PHYSICAL_PAGE_UNAVAILABLE" in out.review_reasons
    assert "text" not in out.as_dict()["tables"][0]["cells"][0]
    assert "text" in out.as_dict(include_text=True)["tables"][0]["cells"][0]
    assert out == inspect_docx_native(path)


def test_docx_native_nested_table_flagged(tmp_path: Path):
    path = tmp_path / "nested.docx"
    doc = Document()
    tab = doc.add_table(rows=1, cols=1)
    tab.cell(0, 0).add_table(rows=1, cols=1)
    doc.save(path)
    table, = inspect_docx_native(path).tables
    assert table.cells[0].nested_table
    assert "NESTED_TABLE_REQUIRES_REVIEW" in table.review_reasons


def test_docx_native_budgets_fail_without_partial_result(tmp_path: Path):
    path = tmp_path / "budget.docx"
    doc = Document()
    tab = doc.add_table(rows=2, cols=2)
    tab.cell(0, 0).text = "shall load"
    doc.save(path)
    with pytest.raises(ValueError, match="cell budget"):
        inspect_docx_native(path, max_cells=1)
    with pytest.raises(ValueError, match="text budget"):
        inspect_docx_native(path, max_total_chars=2)
    with pytest.raises(ValueError, match="size budget"):
        inspect_docx_native(path, max_file_bytes=100)


def test_pdf_native_line_boxes_are_page_bound_without_table_invention(tmp_path: Path):
    path = tmp_path / "source.pdf"
    doc = fitz.open()
    for _ in range(2):
        page = doc.new_page()
        page.insert_text((72, 88), "BT Device ID shall be checked")
    doc.save(str(path))
    doc.close()
    out = inspect_pdf_native(path)
    assert len(out.lines) == 2
    assert [v.page_number for v in out.lines] == [1, 2]
    assert out.lines[0].bbox[0] == 72
    assert out.lines[0].text.startswith("BT Device ID")
    assert "PDF_TABLE_GRID_NOT_VERIFIED" in out.review_reasons
    assert "text" not in out.as_dict()["lines"][0]
    with pytest.raises(ValueError, match="page budget"):
        inspect_pdf_native(path, max_pages=1)
    with pytest.raises(ValueError, match="line budget"):
        inspect_pdf_native(path, max_lines=1)
    with pytest.raises(ValueError, match="text budget"):
        inspect_pdf_native(path, max_total_chars=3)
    assert inspect_native_file(path) == out


def test_unknown_format_rejected(tmp_path: Path):
    path = tmp_path / "source.txt"
    path.write_text("local", encoding="utf-8")
    with pytest.raises(ValueError, match="only DOCX or PDF"):
        inspect_native_file(path)


def test_pdf_opt_in_table_geometry_row_and_column_coordinates(tmp_path: Path):
    path = tmp_path / "grid.pdf"
    doc = fitz.open()
    page = doc.new_page()
    for y in (80, 110, 140):
        page.draw_line((70, y), (370, y))
    for x in (70, 220, 370):
        page.draw_line((x, 80), (x, 140))
    page.insert_text((80, 99), "Condition")
    page.insert_text((230, 99), "Action")
    page.insert_text((80, 129), "NA")
    page.insert_text((230, 129), "Shall load")
    doc.save(str(path))
    doc.close()
    regular = inspect_pdf_native(path)
    assert not regular.pdf_table_candidates  # expensive table search is opt-in
    report = inspect_pdf_native(path, include_table_geometry=True)
    assert len(report.pdf_table_candidates) == 1
    table, = report.pdf_table_candidates
    assert table.page_number == 1 and table.row_count == table.column_count == 2
    assert [(c.row, c.column) for c in table.cells] == [(0,0),(0,1),(1,0),(1,1)]
    assert all(c.status == "GEOMETRY_ONLY_UNVERIFIED" for c in table.cells)
    assert all(c.bbox is not None for c in table.cells)
    assert "PDF_TABLE_SEMANTICS_NOT_VERIFIED" in table.review_reasons
    with pytest.raises(ValueError, match="cell budget"):
        inspect_pdf_native(path, include_table_geometry=True, max_table_cells=3)


def test_japanese_conditional_source_span_remains_exact():
    unit = _source("BT接続の場合、HUはユーザ情報を保存すること。")
    atom, = extract_atomic_candidates(unit)
    evidence = extract_requirement_scope(unit, atom)
    assert evidence.status == "EXPLICIT_PREFIX_CANDIDATE"
    assert evidence.condition.exact_text == "BT接続の場合、"
    assert evidence.action.exact_text == "HUはユーザ情報を保存すること。"


def test_native_cli_defaults_to_no_source_text_and_refuses_overwrite(tmp_path: Path):
    from scripts.inspect_stage1_native_layout import main
    path = tmp_path / "preview.pdf"
    report_path = tmp_path / "private-review.json"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((70, 90), "Shall not disclose engineering secrets")
    pdf.save(str(path))
    pdf.close()
    assert main(["--file", str(path), "--output", str(report_path)]) == 0
    payload = report_path.read_text(encoding="utf-8")
    assert '"source_size"' in payload
    assert "Shall not disclose engineering secrets" not in payload
    with pytest.raises(FileExistsError):
        main(["--file", str(path), "--output", str(report_path)])


def test_local_scope_inspector_reports_counts_without_text(tmp_path: Path, capsys):
    from scripts.inspect_stage1_source import main
    path = tmp_path / "requirements.txt"
    path.write_text("When BT connects, ECU shall notify Body ECU.", encoding="utf-8")
    assert main(["--file", str(path), "--document-id", "source", "--content-id", "1",
                 "--section", "4.1", "--scope"]) == 0
    output = capsys.readouterr().out
    assert '"explicitly_unresolved": 0' in output
    assert '"condition_spans": 1' in output
    assert "When BT connects" not in output
