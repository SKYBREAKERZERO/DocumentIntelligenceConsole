"""v1.14 Stage1 audit UI services, physical joins and real-gold evaluator contracts."""
from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
from pathlib import Path

import pytest
from docx import Document

from app.engineering_intelligence.native_layout import inspect_docx_native, NativeLayoutReport, PdfTableCandidate, PdfGeometryCell
from app.engineering_intelligence.native_review_bundle import native_bundle_summary, prepare_native_table_reviews
from app.engineering_intelligence.review_queue import list_review_queue,load_review_detail,validate_review_scope
from app.engineering_intelligence.source_binding import SourceContent, bind_native_to_sql
from app.engineering_intelligence.stage1_evaluation import evaluate_gold_cases
from app.engineering_intelligence.stage1_review import REVIEW_VERSION,sha


class Cursor:
    def __init__(self,*responses):
        self.responses=list(responses)
        self.calls=[]
        self.last=None
    def execute(self,q,params=None):
        self.calls.append((' '.join(q.split()),params))
        self.last=self.responses.pop(0) if self.responses else []
    def fetchone(self):
        if isinstance(self.last,list):
            return self.last.pop(0) if self.last else None
        return self.last
    def fetchall(self):
        x=self.last
        self.last=[]
        return x if isinstance(x,list) else []


def summary_row(key='a'*64,status=None,event_id=None):
    p=proposal_row(key,status,event_id)
    return (p[0],p[1],p[2],'4.2',['SOURCE_REVIEW_REQUIRED'],status,event_id)


def proposal_row(key='a'*64, status=None, event_id=None):
    payload={'section_id':'4.2','review_reasons':['SOURCE_REVIEW_REQUIRED'],'action':{'exact_text':'ECU shall notify.'}}
    return (key,'SCOPE',sha(payload),payload,status,event_id)


def test_review_queue_scoped_and_receipt_complete():
    cur=Cursor([("21MM",)],[(1,REVIEW_VERSION)],[(1,)],[summary_row()])
    result=list_review_queue(cur,project='21MM',version_id=42,limit=50)
    assert len(result)==1 and result[0].status=='PENDING'
    assert all('s.project_code=%s' in sql for sql,_ in cur.calls if 'JOIN public.specifications s' in sql)
    assert cur.calls[-1][1][0:2]==(42,'21MM')
    assert 'LIMIT %s' in cur.calls[-1][0]


def test_review_queue_digest_and_scope_protection():
    with pytest.raises(ValueError,match='explicit'):
        validate_review_scope('QABUG',1)
    with pytest.raises(ValueError,match='page size'):
        validate_review_scope('21MM',1,1000000)
    with pytest.raises(ValueError,match='cursor'):
        list_review_queue(Cursor(),project='21MM',version_id=1,after_key='xxx')
    altered=list(proposal_row()); altered[2]='g'*64
    with pytest.raises(ValueError,match='metadata identity'):
        list_review_queue(Cursor([("21MM",)],[(1,REVIEW_VERSION)],[(1,)],[tuple(altered[:3]+['4.2',['SOURCE_REVIEW_REQUIRED']]+altered[4:])]),project='21MM',version_id=42)
    with pytest.raises(ValueError,match='incomplete'):
        list_review_queue(Cursor([("21MM",)],[(9,REVIEW_VERSION)],[(2,)]),project='21MM',version_id=42)


def test_review_detail_includes_human_history_and_conflict_protection():
    row=proposal_row(status='CONFIRMED',event_id=7)
    event=(7,'reviewer-1','CONFIRMED','ACTION_ONLY',None,'Reviewed original document','2026-10-14')
    cur=Cursor([("24MM",)],[(1,REVIEW_VERSION)],[(1,)],[row],[event])
    detail=load_review_detail(cur,project='24MM',version_id=12,proposal_key='a'*64)
    assert detail.item.last_event_id==7
    assert detail.history[0]['reason']=='Reviewed original document'
    assert detail.payload['action']['exact_text']=='ECU shall notify.'
    with pytest.raises(ValueError,match='snapshot'):
        load_review_detail(Cursor([("24MM",)],[(1,REVIEW_VERSION)],[(1,)],[row],[]),project='24MM',version_id=12,proposal_key='a'*64)


def test_docx_table_source_bundle_never_confirms_header(tmp_path:Path):
    p=tmp_path/'source.docx'
    d=Document();table=d.add_table(rows=2,cols=2)
    words=['Vehicle profile selection condition','Bluetooth Device ID shall be checked',
           'Region NA exception requirement','Body ECU shall notify the unit']
    for cell,text in zip([c for r in table.rows for c in r.cells],words):cell.text=text
    d.save(p)
    report=inspect_docx_native(p)
    contents=tuple(SourceContent(i+1,'doc',word) for i,word in enumerate(words))
    matches=bind_native_to_sql(report,document_id='doc',expected_source_file_sha256=report.file_sha256,contents=contents)
    bundle=native_bundle_summary(report,matches)
    assert bundle['review_only'] and bundle['table_count']==1
    t=bundle['tables'][0]
    assert t['source_coverage']=='REVIEW_ONLY'
    assert len(t['exact_cell_links'])==4
    assert t['unbound_cells']==0
    assert 'TABLE_HEADER_NOT_AUTOMATICALLY_VERIFIED' in t['structural_flags']
    assert all('source_text_sha256' in c and 'content_id' in c for c in t['exact_cell_links'])
    assert all('text' not in c for c in t['exact_cell_links'])
    with pytest.raises(ValueError,match='drift'):
        bad=(replace(matches[0],text_hash='0'*64),)+matches[1:]
        prepare_native_table_reviews(report,bad)


def test_native_table_missing_cell_is_unverified(tmp_path:Path):
    p=tmp_path/'source.docx';d=Document();t=d.add_table(rows=1,cols=1)
    t.cell(0,0).text='Bluetooth Device ID shall be checked'
    d.save(p)
    report=inspect_docx_native(p)
    matches=bind_native_to_sql(report,document_id='doc',expected_source_file_sha256=report.file_sha256,
                               contents=(SourceContent(1,'doc','This does not match.'),))
    summary=native_bundle_summary(report,matches)
    assert summary['tables'][0]['unbound_cells']==1
    assert not summary['tables'][0]['exact_cell_links']
    assert 'UNBOUND_OR_UNSAFE_NATIVE_CELLS' in summary['tables'][0]['structural_flags']


def test_gold_eval_requires_real_immutable_labels_and_reports_exact_spans():
    text='When Bluetooth connects, the ECU shall notify.'
    case={'document_id':'508-sample','content_id':10,'section_id':'4.2',
          'text':text,'source_sha256':sha256(text.encode()).hexdigest(),
          'expected_requirement_spans':[[0,len(text)]]}
    data={'schema':'stage1-gold-v1','cases':[case]}
    out=evaluate_gold_cases(data)
    assert (out.sample_count,out.true_positives,out.false_positives,out.false_negatives)==(1,1,0,0)
    wrong={**case,'expected_requirement_spans':[[0,4]]}
    out=evaluate_gold_cases({'schema':'stage1-gold-v1','cases':[wrong]})
    assert (out.precision,out.recall)==(0,0)
    with pytest.raises(ValueError,match='nonempty'):
        evaluate_gold_cases({'schema':'stage1-gold-v1','cases':[]})
    with pytest.raises(ValueError,match='checksum'):
        evaluate_gold_cases({'schema':'stage1-gold-v1','cases':[{**case,'source_sha256':'0'*64}]})
    with pytest.raises(ValueError,match='Duplicate gold'):
        evaluate_gold_cases({'schema':'stage1-gold-v1','cases':[case,case]})


def test_main_gui_links_review_window_as_separate_process_without_backend_http():
    from pathlib import Path
    root=Path(__file__).resolve().parents[1]
    src=(root/'app/gui/application.py').read_text(encoding='utf-8')
    assert 'command=self._open_stage1_review_window' in src
    assert 'app.gui.stage1_review_window' in src
    assert 'Stage 1 · 工程证据人工审核' in src
    assert 'record_review_event' not in (root/'app/server/api.py').read_text(encoding='utf-8')


def test_pdf_geometry_budget_is_bounded():
    from app.engineering_intelligence.native_layout import PdfTextLine
    from app.engineering_intelligence.source_binding import native_anchors, SourceBinding
    text='Vehicle ECU shall update local profile after connection.'
    report=NativeLayoutReport('test','PDF','a'*64,100,
        lines=tuple(PdfTextLine(1,0,i,(3.,3.,30.,14.),text+str(i),
                      sha256((text+str(i)).encode()).hexdigest()) for i in range(20)),
        pdf_table_candidates=(PdfTableCandidate(1,0,(0.,0.,40.,40.),1,1,
            (PdfGeometryCell(0,0,(0.,0.,40.,40.)),)),))
    bindings=tuple(SourceBinding(a.anchor_key,a.kind,a.position,a.text_hash,'NO_VERBATIM_SQL_MATCH',
                                 None,None,None,None,('review',)) for a in native_anchors(report))
    with pytest.raises(ValueError,match='geometry comparison budget'):
        prepare_native_table_reviews(report,bindings,max_geometry_checks=10)


def test_database_identity_and_status_filter_cannot_cross_project():
    with pytest.raises(ValueError,match='project identity'):
        list_review_queue(Cursor([('24MM',)]),project='21MM',version_id=42)
    with pytest.raises(ValueError,match='status filter'):
        list_review_queue(Cursor(),project='21MM',version_id=42,status='APPROVED')
    cur=Cursor([('21MM',)],[(1,REVIEW_VERSION)],[(1,)],[summary_row(status='CONFIRMED',event_id=9)])
    rows=list_review_queue(cur,project='21MM',version_id=42,status='CONFIRMED')
    assert rows[0].last_event_id==9
    assert cur.calls[-1][1][4:6]==('CONFIRMED','CONFIRMED')


def test_review_list_uses_only_json_metadata_not_full_proprietary_payload():
    cur=Cursor([('21MM',)],[(1,REVIEW_VERSION)],[(1,)],[summary_row()])
    list_review_queue(cur,project='21MM',version_id=42)
    statement=cur.calls[-1][0]
    assert "p.payload->>'section_id'" in statement
    assert "p.payload->'review_reasons'" in statement
    assert 'p.payload,e.status' not in statement


def test_detail_blocks_tampered_private_evidence_payload():
    row=list(proposal_row()); row[2]='0'*64
    with pytest.raises(ValueError,match='digest mismatch'):
        load_review_detail(Cursor([('21MM',)],[(1,REVIEW_VERSION)],[(1,)],[tuple(row)]),
                           project='21MM',version_id=42,proposal_key='a'*64)


def test_gold_eval_refuses_false_perfect_score_for_unlabelled_dataset():
    text='This section is background only.'
    case={'document_id':'doc','content_id':1,'section_id':'1.1','text':text,
          'source_sha256':sha256(text.encode()).hexdigest(),
          'expected_requirement_spans':[]}
    with pytest.raises(ValueError,match='positively labelled'):
        evaluate_gold_cases({'schema':'stage1-gold-v1','cases':[case]})
