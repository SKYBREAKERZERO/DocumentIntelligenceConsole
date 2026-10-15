"""Review evidence joins must never turn physical geometry into rule approvals."""
from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
from pathlib import Path

import pytest
from docx import Document

from app.engineering_intelligence.native_layout import inspect_docx_native
from app.engineering_intelligence.review_native_compare import compare_review_to_native
from app.engineering_intelligence.review_queue import ReviewDetail, ReviewQueueItem
from app.engineering_intelligence.source_binding import SourceContent, bind_native_to_sql
from app.engineering_intelligence.stage1_review import sha


def _example(tmp_path: Path):
    path = tmp_path / 'evidence.docx'
    doc=Document()
    table=doc.add_table(rows=2,cols=2)
    cells=[c for r in table.rows for c in r.cells]
    values=['Bluetooth Condition', 'Device ID Priority', 'Region NA Exception', 'Body ECU shall notify']
    for cell,value in zip(cells,values):
        cell.text=value
    doc.save(path)
    report=inspect_docx_native(path)
    body=' | '.join(values)
    source=SourceContent(42,'doc508',body)
    matches=bind_native_to_sql(report,document_id='doc508',
             expected_source_file_sha256=report.file_sha256,contents=(source,))
    rows=[]
    for i in range(2):
        row_cells=[]
        for j in range(2):
            phrase=values[i*2+j]
            start=body.index(phrase)
            row_cells.append({'column_index':j,'start_in_content':start,
                              'end_in_content':start+len(phrase),'source_text':phrase})
        rows.append({'unit_key': 'unit'+str(i),'row_kind': 'UNKNOWN','cells':row_cells})
    payload={'document_id':'doc508','content_id':42,'rows':rows,
             'review_reasons':['MERGED_CELL_GEOMETRY_UNAVAILABLE']}
    item=ReviewQueueItem('b'*64,'TABLE',sha(payload),'PENDING',0,None,())
    return report,(source,),matches,ReviewDetail(item,payload,())


def test_exact_docx_cell_alignment_is_always_review_only(tmp_path:Path):
    report,contents,bindings,detail=_example(tmp_path)
    comparison=compare_review_to_native(detail,report,contents,bindings)
    assert comparison.status=='EXACT_NATIVE_CELLS_CANDIDATE'
    assert comparison.exact_cells==comparison.reviewed_cells==4
    assert comparison.human_review_required
    assert all(m['status']=='EXACT_NATIVE_SUBSPAN_CANDIDATE' for m in comparison.matches)
    assert 'TABLE_HEADER_NOT_AUTOMATICALLY_VERIFIED' in comparison.review_flags
    assert 'Bluetooth Condition' in comparison.native_text
    assert not hasattr(comparison,'approved')


def test_missing_or_ambiguous_evidence_never_claimed_as_exact(tmp_path:Path):
    report,contents,bindings,detail=_example(tmp_path)
    altered=tuple(replace(b,status='NO_VERBATIM_SQL_MATCH',content_id=None,
                          start_offset=None,end_offset=None,requirement_unit_key=None)
                  if b.kind=='DOCX_CELL' and b.position[1:]==(0,0) else b for b in bindings)
    comparison=compare_review_to_native(detail,report,contents,altered)
    assert comparison.status=='PARTIAL_OR_AMBIGUOUS'
    assert comparison.exact_cells==3
    assert 'NO_EXACT_NATIVE_CELL' in comparison.review_flags


def test_proposal_digest_and_offsets_are_independently_verified(tmp_path:Path):
    report,contents,bindings,detail=_example(tmp_path)
    with pytest.raises(ValueError,match='digest'):
        compare_review_to_native(replace(detail,payload={**detail.payload,'content_id':43}),
                                 report,contents,bindings)
    tampered={**detail.payload, 'rows':[dict(x) for x in detail.payload['rows']]}
    tampered['rows'][0]['cells']=[dict(x) for x in tampered['rows'][0]['cells']]
    tampered['rows'][0]['cells'][0]['source_text']='wrong'
    altered=replace(detail,payload=tampered,item=replace(detail.item,digest=sha(tampered)))
    with pytest.raises(ValueError,match='does not match'):
        compare_review_to_native(altered,report,contents,bindings)


def test_incomplete_table_source_never_promoted(tmp_path:Path):
    report,contents,bindings,detail=_example(tmp_path)
    modified=tuple(SourceContent(42,'doc508','NO LONGER THE SAME SQL TEXT') for _ in contents)
    with pytest.raises(ValueError,match='does not match'):
        compare_review_to_native(detail,report,modified,bindings)


def test_native_alignment_bounded(tmp_path:Path):
    report,contents,bindings,detail=_example(tmp_path)
    with pytest.raises(ValueError,match='budget'):
        compare_review_to_native(detail,report,contents,bindings,max_review_cells=1)


def test_gui_source_compare_is_local_readonly_and_not_network_approval():
    src=(Path(__file__).resolve().parents[1]/'app/gui/stage1_review_window.py').read_text(encoding='utf-8')
    assert 'askopenfilename' in src
    assert 'read_verified_native_context' in src
    assert 'READ ONLY' in src
    assert 'compare_review_to_native' in src
    assert 'record_review_event' in src  # separate human action remains
    assert 'source_path=path' in src


def test_scope_can_map_docx_body_paragraph_and_never_auto_confirm(tmp_path:Path):
    p=tmp_path/'plain.docx'
    phrase='When Bluetooth is paired the vehicle shall load the profile.'
    d=Document();d.add_paragraph(phrase);d.save(p)
    report=inspect_docx_native(p)
    source=SourceContent(100,'doc508',phrase)
    bindings=bind_native_to_sql(report,document_id='doc508',
                                expected_source_file_sha256=report.file_sha256,
                                contents=(source,))
    body={'source_content_id':100,'candidate_source_start':5,
          'candidate_source_end':len(phrase),'candidate_text_hash':sha256(phrase[5:].encode()).hexdigest(),
          'review_reasons':['SCOPE_HUMAN_REVIEW_REQUIRED']}
    detail=ReviewDetail(ReviewQueueItem('a'*64,'SCOPE',sha(body),'PENDING',0,None,()),body,())
    result=compare_review_to_native(detail,report,(source,),bindings)
    assert result.exact_cells==1
    assert 'DOCX_PARAGRAPH:' in result.native_text
    assert result.matches[0]['sql_start']==5
    assert result.human_review_required


def test_large_content_preview_windows_to_actual_target(tmp_path:Path):
    report,contents,_,detail=_example(tmp_path)
    source=SourceContent(42,'doc508','A'*50_000 + '\n' + contents[0].text)
    bindings=bind_native_to_sql(report,document_id='doc508',
        expected_source_file_sha256=report.file_sha256,contents=(source,))
    extra=len(source.text)-len(contents[0].text)
    payload=dict(detail.payload)
    rows=[]
    for row in payload['rows']:
        cells=[{**c,'start_in_content':c['start_in_content']+extra,
                'end_in_content':c['end_in_content']+extra} for c in row['cells']]
        rows.append({**row,'cells':cells})
    payload['rows']=rows
    detail=replace(detail,payload=payload,item=replace(detail.item,digest=sha(payload)))
    result=compare_review_to_native(detail,report,(source,),bindings,max_display_chars=800)
    assert result.sql_excerpt_start >= 49_000
    assert result.status=='EXACT_NATIVE_CELLS_CANDIDATE'
    assert all(result.sql_text[m['sql_start']-result.sql_excerpt_start:
               m['sql_end']-result.sql_excerpt_start]
               == source.text[m['sql_start']:m['sql_end']] for m in result.matches)
    with pytest.raises(ValueError,match='display budget'):
        compare_review_to_native(detail,report,(source,),bindings,max_display_chars=5)


class _SourceCursor:
    def __init__(self,identity,ownership):
        self.responses=[[identity],[ownership]]
        self.statements=[]
    def execute(self,sql,params=None):
        self.statements.append((' '.join(sql.split()),params))
    def fetchone(self):
        return self.responses.pop(0)[0]


def test_source_provenance_validation_precedes_large_sql_fetch(tmp_path:Path):
    from app.engineering_intelligence.review_native_compare import _source_rows_for_version
    report,_,_,_=_example(tmp_path)
    c=_SourceCursor(('21MM',),('doc508','0'*64))
    with pytest.raises(ValueError,match='SHA-256'):
        _source_rows_for_version(c,project='21MM',version_id=12,report=report,
                                 max_contents=100,max_source_chars=10000,max_units=100)
    assert not any('FROM public.contents' in q for q,_ in c.statements)


def test_readonly_cli_offers_proposal_level_native_alignment_without_dumping_text():
    src=(Path(__file__).resolve().parents[1]/'scripts/inspect_stage1_sql_binding.py').read_text(encoding='utf-8')
    assert '--review-proposal-key' in src
    assert 'load_review_detail(' in src
    assert 'compare_review_to_native' in src
    assert "'exact_cells': comparison.exact_cells" in src
    assert "'sql_text': comparison.sql_text" not in src
    assert "'native_text': comparison.native_text" not in src
