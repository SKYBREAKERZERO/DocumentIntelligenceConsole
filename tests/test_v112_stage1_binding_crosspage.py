"""Stage 1 v1.12: synthetic, local, provenance-first integration regression."""
from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
from pathlib import Path

import fitz
import pytest
from docx import Document

from app.engineering_intelligence.crosspage import suggest_crosspage_tables
from app.engineering_intelligence.native_layout import (
    PdfGeometryCell, PdfTableCandidate, NativeLayoutReport,
    inspect_docx_native, inspect_pdf_native,
)
from app.engineering_intelligence.requirements import extract_requirement_units
from app.engineering_intelligence.source_binding import (
    SourceContent, bind_native_to_sql, native_anchors,
)


def _hash(text: str) -> str:
    return sha256(text.encode()).hexdigest()


def _docx_report(tmp_path: Path, text: str):
    p = tmp_path / 'r.docx'
    doc = Document()
    tab = doc.add_table(rows=1, cols=1)
    tab.cell(0, 0).text = text
    doc.save(p)
    return inspect_docx_native(p)


def _binding(report, texts, units=(), **kwargs):
    rows = tuple(SourceContent(i + 1, 'doc', v) for i, v in enumerate(texts))
    return bind_native_to_sql(
        report, document_id='doc', expected_source_file_sha256=report.file_sha256,
        contents=rows, units=units, **kwargs,
    )


def test_native_exact_unique_content_and_requirement(tmp_path):
    text = 'ECU shall notify Body ECU within five seconds.'
    report = _docx_report(tmp_path, text)
    body = 'Prefix. ' + text + ' More text.'
    unit = extract_requirement_units(document_id='doc', content_id=1, section_id='4.1', text=body)
    matches = _binding(report, [body], unit)
    assert len(matches) == 1
    result = matches[0]
    assert result.content_id == 1
    assert body[result.start_offset:result.end_offset] == text
    # Only assign a Requirement when the entire native anchor lies inside exactly one unit.
    assert result.status in ('UNIQUE_CONTENT_SPAN_CANDIDATE','UNIQUE_REQUIREMENT_SPAN_CANDIDATE')
    assert result.review_reasons
    assert result.text_hash == _hash(text)


def test_duplicate_sql_text_refuses_unique_binding(tmp_path):
    text = 'ECU shall notify Body ECU after connection.'
    report = _docx_report(tmp_path, text)
    matches = _binding(report, [text, text])
    assert matches[0].status == 'AMBIGUOUS_VERBATIM_SQL_MATCH'
    assert matches[0].content_id is None


def test_repeated_text_inside_one_chunk_refuses_unique(tmp_path):
    text = 'The vehicle shall save the active profile.'
    report = _docx_report(tmp_path, text)
    matches = _binding(report, [text + ' ' + text])
    assert matches[0].status == 'AMBIGUOUS_VERBATIM_SQL_MATCH'


def test_missing_and_short_evidence_are_explicit(tmp_path):
    report = _docx_report(tmp_path, 'NA')
    result, = _binding(report, ['NA'])
    assert result.status == 'TOO_SHORT_FOR_UNIQUE_BINDING'
    report2 = _docx_report(tmp_path, 'The HU shall update Memory ID.')
    result2, = _binding(report2, ['The HU will update Memory ID.'])
    assert result2.status == 'NO_VERBATIM_SQL_MATCH'


def test_untrusted_source_file_rejected_before_text_matching(tmp_path):
    report = _docx_report(tmp_path, 'The HU shall update Memory ID.')
    with pytest.raises(ValueError, match='source SHA-256'):
        bind_native_to_sql(report, document_id='doc', expected_source_file_sha256='0' * 64,
                           contents=(SourceContent(1,'doc','The HU shall update Memory ID.'),))
    with pytest.raises(ValueError, match='missing or mismatched'):
        bind_native_to_sql(report, document_id='doc', expected_source_file_sha256='', contents=())


def test_cross_project_source_and_requirement_drift_rejected(tmp_path):
    report = _docx_report(tmp_path, 'ECU shall notify Body ECU after connection.')
    with pytest.raises(ValueError, match='Foreign'):
        bind_native_to_sql(report, document_id='doc', expected_source_file_sha256=report.file_sha256,
                           contents=(SourceContent(1, 'foreign','ECU shall notify Body ECU after connection.'),))
    unit, = extract_requirement_units(document_id='doc', content_id=1,
                                      section_id='4.1', text='ECU shall notify Body ECU after connection.')
    with pytest.raises(ValueError, match='source span drift'):
        _binding(report, ['ECU shall notify Body ECU after connection.'], (replace(unit, text='injected'),))


def test_native_cells_are_deduplicated_for_vertical_merges(tmp_path):
    p = tmp_path / 'merged.docx'
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0,0).text = 'Long technical condition that must hold.'
    table.cell(0,0).merge(table.cell(1,0))
    doc.save(p)
    report = inspect_docx_native(p)
    anchors = native_anchors(report)
    assert len(anchors) == 1


def test_source_index_budget_fail_closed(tmp_path):
    report = _docx_report(tmp_path, 'ECU shall notify the body controller immediately.')
    with pytest.raises(ValueError, match='text budget'):
        _binding(report, ['some huge source text'], max_source_chars=3)
    with pytest.raises(ValueError, match='budgets'):
        _binding(report, ['sample'], max_anchors=0)
    with pytest.raises(ValueError, match='index budget'):
        _binding(report, ['sample'], max_automaton_nodes=2)


def _table(page: int, idx: int, *, top: float, bottom: float, x: tuple[float, ...],
           height: float = 800) -> PdfTableCandidate:
    cells = tuple(PdfGeometryCell(0, col, (x[col], top, x[col+1], bottom)) for col in range(len(x)-1))
    return PdfTableCandidate(page, idx, (x[0], top, x[-1], bottom), 1, len(x)-1, cells,
                             page_width=600., page_height=height)


def _pdf_tables(*tables):
    return NativeLayoutReport('test', 'PDF', 'a'*64, 100, pdf_table_candidates=tuple(tables))


def test_pdf_adjacent_page_geometry_generates_review_only_continuation():
    report = _pdf_tables(
        _table(1, 1, top=650, bottom=760, x=(70,180,350)),
        _table(2, 2, top=50, bottom=200, x=(70,180,350)),
    )
    result, = suggest_crosspage_tables(report)
    assert (result.previous_page, result.next_page) == (1,2)
    assert result.status == 'POSSIBLE_CONTINUATION_REVIEW_ONLY'
    assert 'HEADER_AND_SEMANTIC_CONTINUITY_UNVERIFIED' in result.reasons


@pytest.mark.parametrize('next_table', [
    _table(2,2,top=120,bottom=200,x=(70,180,350)),
    _table(2,2,top=50,bottom=200,x=(100,210,380)),
    _table(2,2,top=50,bottom=200,x=(70,350)),
    _table(3,2,top=50,bottom=200,x=(70,180,350)),
])
def test_unrelated_pdf_tables_not_linked(next_table):
    report = _pdf_tables(_table(1,1,top=650,bottom=760,x=(70,180,350)), next_table)
    assert not suggest_crosspage_tables(report)


def test_crosspage_budget_blocks_unbounded_cartesian_work():
    report = _pdf_tables(_table(1,1,top=690,bottom=770,x=(70,180,350)),
                         _table(2,2,top=50,bottom=140,x=(70,180,350)))
    with pytest.raises(ValueError, match='budget'):
        suggest_crosspage_tables(report, max_pairs=0)


def test_native_pdf_line_preview_does_not_write_sql(tmp_path):
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((70,90),'ECU shall perform notification after connection.')
    path = tmp_path / 'local.pdf'
    pdf.save(str(path))
    pdf.close()
    report = inspect_pdf_native(path)
    text = report.lines[0].text
    candidate, = _binding(report, ['before ' + text + ' after'])
    assert candidate.status == 'UNIQUE_CONTENT_SPAN_CANDIDATE'
    assert candidate.kind == 'PDF_LINE'
    assert candidate.position[0] == 1


def test_sql_pdf_page_mismatch_rejected_without_binding(tmp_path):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((70,90),'ECU shall notify the user after Bluetooth connection.')
    path = tmp_path/'page.pdf'
    doc.save(str(path)); doc.close()
    report = inspect_pdf_native(path)
    value = report.lines[0].text
    result, = bind_native_to_sql(report, document_id='doc',
        expected_source_file_sha256=report.file_sha256,
        contents=(SourceContent(1,'doc',value,2),))
    assert result.status == 'PAGE_NUMBER_MISMATCH'
    assert result.content_id is None
    accepted, = bind_native_to_sql(report, document_id='doc',
        expected_source_file_sha256=report.file_sha256,
        contents=(SourceContent(1,'doc',value,1),))
    assert accepted.status == 'UNIQUE_CONTENT_SPAN_CANDIDATE'
    assert 'SQL_PAGE_UNAVAILABLE' not in accepted.review_reasons


def test_two_native_cells_cannot_claim_same_sql_location(tmp_path):
    p = tmp_path / 'two.docx'
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    for row in table.rows:
        row.cells[0].text = 'ECU shall update profile after Bluetooth connection.'
    doc.save(p)
    report = inspect_docx_native(p)
    bindings = _binding(report, ['ECU shall update profile after Bluetooth connection.'])
    assert len(bindings) == 2
    assert all(b.status == 'AMBIGUOUS_NATIVE_SOURCE_ANCHORS' for b in bindings)
    assert all(b.content_id is None for b in bindings)


def test_large_but_bounded_unique_native_index():
    from app.engineering_intelligence.native_layout import PdfTextLine
    count = 1200
    texts = [f'Exact system rule {i:05d} shall maintain original evidence.' for i in range(count)]
    report = NativeLayoutReport('test', 'PDF', 'a'*64, 123,
        lines=tuple(PdfTextLine(i+1,0,0,(1.,2.,3.,4.),text,_hash(text))
                    for i,text in enumerate(texts)))
    contents = tuple(SourceContent(i+1,'doc',text,i+1) for i,text in enumerate(texts))
    results = bind_native_to_sql(report, document_id='doc', expected_source_file_sha256='a'*64,
                                 contents=contents, max_anchors=2000, max_source_chars=200000)
    assert len(results) == count
    assert all(b.status == 'UNIQUE_CONTENT_SPAN_CANDIDATE' for b in results)


def test_docx_native_body_paragraphs_have_stable_order_and_exact_sql_binding(tmp_path):
    path = tmp_path / 'body.docx'
    doc = Document()
    doc.add_paragraph('When Bluetooth is connected, the HU shall load profile.')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0,0).text = 'The ECU shall notify the connected device.'
    doc.add_paragraph('Unless NA applies, the HU must load local settings.')
    doc.save(path)
    report = inspect_docx_native(path)
    assert [p.body_order for p in report.paragraphs] == [0,2]
    assert len(native_anchors(report)) == 3
    assert 'text' not in report.as_dict()['paragraphs'][0]
    assert report.as_dict(include_text=True)['paragraphs'][0]['text'].startswith('When Bluetooth')
    values = [p.text for p in report.paragraphs] + [report.tables[0].cells[0].text]
    bindings = _binding(report, values)
    assert len(bindings) == 3
    assert all(b.status == 'UNIQUE_CONTENT_SPAN_CANDIDATE' for b in bindings)
    assert [b.kind for b in bindings].count('DOCX_PARAGRAPH') == 2
    with pytest.raises(ValueError, match='paragraph budget'):
        inspect_docx_native(path, max_paragraphs=1)


def test_sql_binding_cli_rejects_wrong_authoritative_file_before_source_export(tmp_path, monkeypatch):
    import sys
    from types import SimpleNamespace
    from scripts.inspect_stage1_sql_binding import main
    p = tmp_path / 'unauthorized.pdf'
    pdf = fitz.open()
    pdf.new_page().insert_text((70,90),'Vehicle should select the correct profile.')
    pdf.save(str(p)); pdf.close()
    statements = []

    class Cursor:
        def __enter__(self): return self
        def __exit__(self,*exc): return False
        def execute(self, sql, args=None):
            statements.append(sql)
        def fetchone(self):
            query = statements[-1]
            if 'project_info' in query:
                return ('21MM',)
            if 'specification_versions sv' in query:
                return ('doc-1','0'*64)  # authoritative file differs
            raise AssertionError('Source contents must not be queried on failed provenance')
    class Connection:
        def __enter__(self): return self
        def __exit__(self,*exc): return False
        def cursor(self): return Cursor()
    monkeypatch.setitem(sys.modules, 'psycopg', SimpleNamespace(connect=lambda **kw: Connection()))
    monkeypatch.setenv('POSTGRES_PASSWORD','test-only')
    with pytest.raises(ValueError, match='source_hash'):
        main(['--file',str(p),'--project','21MM','--database','rag_21mm','--version-id','12'])
    assert statements and 'READ ONLY' in statements[0]
    assert not any('INSERT' in query or 'UPDATE' in query or 'DELETE' in query for query in statements)
