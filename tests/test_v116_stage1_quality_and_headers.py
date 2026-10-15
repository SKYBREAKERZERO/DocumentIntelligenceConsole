"""v1.16 deterministic SYNTHETIC unit tests, not real 508/550 acceptance."""
from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document

from app.engineering_intelligence.native_layout import inspect_docx_native
from app.engineering_intelligence.source_binding import SourceContent, bind_native_to_sql
from app.engineering_intelligence.table_header_review import (
    preview_table_header_columns, redacted_header_review)
from app.engineering_intelligence.stage1_evaluation import (
    GOLD_SCHEMA_V2, evaluate_gold_quality_cases)
from scripts.evaluate_stage1_gold import main


def _case(text: str, content_id: int, status: str, parts: dict[str, list[int] | None]):
    return {"document_id": "synthetic-not-real-508", "content_id": content_id,
            "section_id": "test", "text": text,
            "source_sha256": hashlib.sha256(text.encode()).hexdigest(),
            "expected_requirement_spans": [[0, len(text)]],
            "scope_annotations": [{"requirement_span": [0, len(text)],
                                   "status": status,
                                   **{f"{k}_span": v for k, v in parts.items()}}]}


def test_gold_v2_separately_measures_exact_scope_and_requirement():
    text = 'When Bluetooth connects, the ECU shall notify.'
    gold = {"schema": GOLD_SCHEMA_V2, "cases": [
        _case(text, 1, "CONDITION_GOVERNS_ACTION", {
            "condition": [0, 24], "action": [24, len(text)]})]}
    stats = evaluate_gold_quality_cases(gold)
    assert stats.requirement_metrics.f1 == 1
    assert stats.scope_f1 == 1
    assert stats.scope_true_positives == 2
    assert stats.status == 'VALIDATED_GOLD'
    assert "text" not in json.dumps(stats.__dict__ if hasattr(stats, '__dict__') else str(stats))


def test_gold_v2_negative_scope_and_status_not_falsely_perfect():
    text = 'When Bluetooth connects, the ECU shall notify.'
    data = {"schema": GOLD_SCHEMA_V2, "cases": [
        _case(text, 1, "UNRESOLVED", {})]}
    res = evaluate_gold_quality_cases(data)
    assert res.requirement_metrics.f1 == 1
    assert res.scope_f1 is None and res.status == 'NO_SCOPE_POSITIVES'
    assert res.scope_false_positives == 2
    assert res.unresolved_gold_requirements == 1


def test_gold_v2_detects_wrong_rule_parts_without_fabricating_answers():
    text = 'Unless region is NA, ECU shall save the profile.'
    data = {"schema": GOLD_SCHEMA_V2, "cases": [
        _case(text, 1, "EXCEPTION_OVERRIDES_ACTION", {
            "exception": [0, 20], "action": [20, len(text)]})]}
    assert evaluate_gold_quality_cases(data).scope_f1 == 1
    bad = _case(text, 1, "EXCEPTION_OVERRIDES_ACTION", {
        "exception": [0, 18], "action": [20, len(text)]})
    mismatch = evaluate_gold_quality_cases({"schema": GOLD_SCHEMA_V2, "cases": [bad]})
    assert mismatch.scope_true_positives == 1
    assert 0 < mismatch.scope_f1 < 1


@pytest.mark.parametrize('mutation,pattern', [
    (lambda x: x.pop('scope_annotations'), 'complete human'),
    (lambda x: x['scope_annotations'][0].update(status='CONFIRMED'), 'Unknown human'),
    (lambda x: x['scope_annotations'][0].update(condition_span=[100,120]), 'outside'),
    (lambda x: x['scope_annotations'][0].update(action_span=[5,40]), 'Overlapping'),
    (lambda x: x['scope_annotations'].append(dict(x['scope_annotations'][0])), 'unique'),
    (lambda x: x['scope_annotations'].clear(), 'Every positive'),
])
def test_gold_v2_fails_closed_for_invalid_label_sets(mutation, pattern):
    text = 'When Bluetooth connects, the ECU shall notify.'
    data = {"schema": GOLD_SCHEMA_V2, "cases": [
        _case(text, 1, "CONDITION_GOVERNS_ACTION", {
            "condition": [0, 24], "action": [24, len(text)]})]}
    mutation(data['cases'][0])
    with pytest.raises(ValueError, match=pattern):
        evaluate_gold_quality_cases(data)


def test_gold_v2_ci_gate_distinguishes_measured_and_missing(tmp_path: Path, capsys):
    text = 'ECU shall notify.'
    data = {"schema": GOLD_SCHEMA_V2, "cases": [
        _case(text, 1, "ACTION_ONLY", {"action": [0, len(text)]})]}
    p = tmp_path / 'gold-v2.json'
    p.write_text(json.dumps(data), encoding='utf-8')
    assert main(['--gold', str(p), '--min-requirement-f1', '0.99', '--min-scope-f1', '0.99']) == 0
    assert 'acceptance_passed' in capsys.readouterr().out
    data['cases'][0]['scope_annotations'][0]['action_span'] = [0, 3]
    p.write_text(json.dumps(data), encoding='utf-8')
    assert main(['--gold', str(p), '--min-scope-f1', '0.99']) == 2
    assert 'false' in capsys.readouterr().out.lower()


def _table(tmp_path: Path, *, merged: bool = False):
    file = tmp_path / 'synthetic.docx'
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    values = [['Bluetooth feature group', 'User profile assignment'], ['ECU shall notify', 'BT Device ID shall match'],
              ['load local profile data', 'load remote profile data']]
    for row, words in zip(table.rows, values):
        for cell, word in zip(row.cells, words):
            cell.text = word
    if merged:
        table.cell(0, 0).merge(table.cell(0, 1))
    doc.save(file)
    report = inspect_docx_native(file)
    texts = [cell.text for cell in report.tables[0].cells if cell.text.strip()]
    contents = tuple(SourceContent(i + 1, 'doc', t) for i, t in enumerate(texts))
    bindings = bind_native_to_sql(report, document_id='doc',
                                  expected_source_file_sha256=report.file_sha256,
                                  contents=contents)
    return report, bindings


def test_explicit_docx_header_evidence_matrix_is_review_only(tmp_path: Path):
    report, binding = _table(tmp_path)
    evidence = preview_table_header_columns(report, binding,
                            table_identity='DOCX:0', header_rows=(0,))
    assert evidence.status == 'REVIEW_ONLY'
    assert evidence.exact_header_cells == 2
    assert evidence.exact_data_cells == 4
    assert len(evidence.columns) == 2
    assert all(len(c.header_links) == 1 and len(c.data_links) == 2 for c in evidence.columns)
    assert 'TABLE_HEADER_NOT_AUTOMATICALLY_VERIFIED' in evidence.review_blockers
    assert 'Bluetooth feature group' not in json.dumps(redacted_header_review(evidence))


def test_header_evidence_keeps_merged_header_unapproved(tmp_path: Path):
    report, binding = _table(tmp_path, merged=True)
    preview = preview_table_header_columns(report, binding,
                             table_identity='DOCX:0', header_rows=(0,))
    assert preview.status == 'REVIEW_ONLY'
    assert 'MERGED_HEADER_OR_DATA_CELL_REQUIRES_HUMAN_REVIEW' in preview.review_blockers
    assert 'SHARED_MERGED_HEADER_ANCHOR' in preview.review_blockers


@pytest.mark.parametrize('rows', [(), (2,), (1,0), (0,0), (-1,), (0,1,2), (True,), (999,)])
def test_header_evidence_rejects_missing_ambiguous_or_invalid_header_rows(tmp_path: Path, rows):
    report, bindings = _table(tmp_path)
    with pytest.raises(ValueError):
        preview_table_header_columns(report, bindings, table_identity='DOCX:0', header_rows=rows)


def test_header_evidence_rejects_binding_drift_and_foreign_identity(tmp_path: Path):
    report, bindings = _table(tmp_path)
    with pytest.raises(ValueError, match='does not exist'):
        preview_table_header_columns(report, bindings,
                                     table_identity='DOCX:200', header_rows=(0,))
    with pytest.raises(ValueError, match='drift'):
        preview_table_header_columns(report,
            (replace(bindings[0], text_hash='0'*64),) + bindings[1:],
            table_identity='DOCX:0', header_rows=(0,))


def test_annotation_scaffold_never_auto_labels_and_refuses_identity_collision():
    from scripts.prepare_stage1_gold_v2 import build_unlabelled_gold
    source = {"schema": "stage1-source-items-v1", "cases": [{
        "document_id": "trusted-local-doc", "content_id": 20,
        "section_id": "4.2", "text": "When connected, ECU shall notify."}]}
    scaffold = build_unlabelled_gold(source)
    assert scaffold['schema'] == GOLD_SCHEMA_V2
    assert scaffold['cases'][0]['expected_requirement_spans'] == []
    assert scaffold['cases'][0]['scope_annotations'] == []
    with pytest.raises(ValueError, match='positively labelled'):
        evaluate_gold_quality_cases(scaffold)
    with pytest.raises(ValueError, match='Duplicate'):
        build_unlabelled_gold({**source, "cases": source['cases'] * 2})
    with pytest.raises(ValueError, match='drift'):
        build_unlabelled_gold({**source, 'cases': [{**source['cases'][0],
            'source_sha256': '0' * 64}]})


def test_annotation_cli_refuses_unapproved_copy_and_overwrite(tmp_path: Path, capsys):
    from scripts.prepare_stage1_gold_v2 import main as prepare_main
    raw = tmp_path / 'source.json'
    out = tmp_path / 'gold.json'
    raw.write_text(json.dumps({'schema': 'stage1-source-items-v1', 'cases': [
        {'document_id': 'doc', 'content_id': 1, 'section_id': '4',
         'text': 'ECU shall notify.'}]}), encoding='utf-8')
    with pytest.raises(SystemExit):
        prepare_main(['--source-json', str(raw), '--out', str(out)])
    assert not out.exists()
    assert prepare_main(['--source-json', str(raw), '--out', str(out),
                        '--confirm-local-sensitive-copy']) == 0
    assert 'UNLABELLED' in capsys.readouterr().out
    with pytest.raises(FileExistsError):
        prepare_main(['--source-json', str(raw), '--out', str(out),
                      '--confirm-local-sensitive-copy'])


def test_v2_disallows_overlapping_gold_requirements():
    text = 'ECU shall notify. ECU shall save.'
    case = _case(text, 1, 'ACTION_ONLY', {'action': [0, len(text)]})
    case['expected_requirement_spans'] = [[0, 20], [10, len(text)]]
    case['scope_annotations'] = []
    with pytest.raises(ValueError, match='Overlapping'):
        evaluate_gold_quality_cases({'schema': GOLD_SCHEMA_V2, 'cases': [case]})


def test_header_evidence_pdf_remains_geometry_only_and_unapproved():
    from app.engineering_intelligence.native_layout import (
        NativeLayoutReport, PdfTableCandidate, PdfGeometryCell, PdfTextLine)
    from app.engineering_intelligence.source_binding import SourceBinding, native_anchors
    texts = ('Region header description', 'Action header description',
             'Japanese region specification', 'ECU shall notify by CAN bus')
    lines = tuple(PdfTextLine(1, i, 0, (10.0 if i % 2 == 0 else 50.,
                10.0 if i < 2 else 30., 45. if i % 2 == 0 else 95.,
                22. if i < 2 else 42.), t, hashlib.sha256(t.encode()).hexdigest())
                  for i, t in enumerate(texts))
    physical = (
        PdfGeometryCell(0, 0, (5., 5., 47., 25.)),
        PdfGeometryCell(0, 1, (48., 5., 98., 25.)),
        PdfGeometryCell(1, 0, (5., 26., 47., 45.)),
        PdfGeometryCell(1, 1, (48., 26., 98., 45.)),
    )
    report = NativeLayoutReport('test', 'PDF', 'f'*64, 100, lines=lines,
       pdf_table_candidates=(PdfTableCandidate(1, 0, (5.,5.,98.,45.), 2, 2, physical),))
    anchors = native_anchors(report)
    bindings = tuple(SourceBinding(a.anchor_key, a.kind, a.position, a.text_hash,
            'UNIQUE_CONTENT_SPAN_CANDIDATE', i+1, 0, len(a.text), None, ())
            for i, a in enumerate(anchors))
    preview = preview_table_header_columns(report, bindings, table_identity='PDF:1:0',
                                          header_rows=(0,))
    assert preview.status == 'REVIEW_ONLY'
    assert preview.exact_header_cells == 2
    assert 'PDF_HEADER_AND_CROSSPAGE_SEMANTICS_UNVERIFIED' in preview.review_blockers


def test_native_header_proposal_cannot_attach_unrelated_table_or_partial_match():
    from types import SimpleNamespace
    from app.engineering_intelligence.table_header_review import require_proposal_matches_header_table
    ok = SimpleNamespace(status='EXACT_NATIVE_CELLS_CANDIDATE',
                         matches=({'native_table': 'DOCX:0',
                                   'status': 'EXACT_NATIVE_SUBSPAN_CANDIDATE'},))
    require_proposal_matches_header_table(ok, 'DOCX:0')
    with pytest.raises(ValueError, match='does not uniquely match'):
        require_proposal_matches_header_table(ok, 'DOCX:1')
    with pytest.raises(ValueError, match='not completely'):
        require_proposal_matches_header_table(SimpleNamespace(status='PARTIAL_OR_AMBIGUOUS',
                                                              matches=ok.matches), 'DOCX:0')
    with pytest.raises(ValueError, match='No matched'):
        require_proposal_matches_header_table(SimpleNamespace(status='EXACT_NATIVE_CELLS_CANDIDATE',
                                                              matches=()), 'DOCX:0')


def test_header_review_column_membership_budget_fails_closed(tmp_path: Path):
    report, bindings = _table(tmp_path)
    with pytest.raises(ValueError, match='membership budget'):
        preview_table_header_columns(report, bindings, table_identity='DOCX:0',
                                     header_rows=(0,), max_column_memberships=1)
