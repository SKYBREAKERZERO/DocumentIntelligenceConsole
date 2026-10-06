from __future__ import annotations

import csv
import hashlib
import json
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator, Sequence

from .canonical import QABugCanonicalizer, canonical_header


class QABugImportError(RuntimeError):
    pass


@dataclass(slots=True)
class QABugImportReport:
    source_file: str
    record_type: str
    rows_seen: int = 0
    rows_valid: int = 0
    rows_skipped: int = 0
    rows_failed: int = 0
    rows_quarantined: int = 0
    records_upserted: int = 0
    chunks_upserted: int = 0
    jsonl_written: int = 0
    qdrant_points_upserted: int = 0
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.rows_failed == 0


class SourceRowValues(list):
    """List-like source row carrying deterministic repair provenance."""

    def __init__(self, values=(), *, repair_info: dict[str, Any] | None = None) -> None:
        super().__init__(values)
        self.repair_info = dict(repair_info or {})


class QABugFileReader:
    """Streaming reader for QA XLSX and Jira BUG CSV/TSV/TXT exports."""

    @staticmethod
    def detect_encoding(path: str | Path) -> str:
        raw = Path(path).read_bytes()[:65536]
        if raw.startswith(b"\xef\xbb\xbf"):
            return "utf-8-sig"
        if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            return "utf-16"
        for encoding in ("utf-8", "cp932", "shift_jis"):
            try:
                raw.decode(encoding)
                return encoding
            except UnicodeDecodeError:
                pass
        return "utf-8"

    @classmethod
    def detect_delimiter(cls, path: str | Path, encoding: str | None = None) -> str:
        enc = encoding or cls.detect_encoding(path)
        with Path(path).open("r", encoding=enc, errors="replace", newline="") as handle:
            sample = handle.read(65536)
        counts = {"\t": sample.count("\t"), ",": sample.count(","), ";": sample.count(";")}
        return max(counts, key=counts.get) if max(counts.values()) else "\t"

    @staticmethod
    def _looks_like_header(values: Sequence[Any]) -> bool:
        recognized = sum(1 for value in values if canonical_header(value))
        has_key = any(canonical_header(value) == "issue_key" for value in values)
        return has_key and recognized >= 5

    @classmethod
    def iter_xlsx(
        cls,
        path: str | Path,
        *,
        sheet_name: str | None = None,
    ) -> Iterator[tuple[str, int, Sequence[Any], Sequence[Any]]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise QABugImportError("openpyxl is required for QA/BUG XLSX import.") from exc
        workbook = load_workbook(filename=Path(path), read_only=True, data_only=True)
        try:
            sheets = [workbook[sheet_name]] if sheet_name else list(workbook.worksheets)
            for sheet in sheets:
                header: tuple[Any, ...] | None = None
                header_row = 0
                for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = tuple(row)
                    if header is None:
                        if row_index <= 30 and cls._looks_like_header(values):
                            header = values
                            header_row = row_index
                        continue
                    if not any(str(value or "").strip() for value in values):
                        continue
                    yield sheet.title, row_index, header, values
                if header is None:
                    continue
        finally:
            workbook.close()

    @classmethod
    def iter_delimited(
        cls,
        path: str | Path,
        *,
        encoding: str | None = None,
        delimiter: str | None = None,
    ) -> Iterator[tuple[str, int, Sequence[Any], Sequence[Any]]]:
        enc = encoding or cls.detect_encoding(path)
        delim = delimiter or cls.detect_delimiter(path, enc)
        csv.field_size_limit(max(csv.field_size_limit(), 128 * 1024 * 1024))
        with Path(path).open("r", encoding=enc, errors="replace", newline="") as handle:
            reader = csv.reader(handle, delimiter=delim)
            try:
                header = next(reader)
            except StopIteration:
                return
            if not cls._looks_like_header(header):
                raise QABugImportError(
                    f"Could not identify Jira/QA header in {Path(path).name}; "
                    "expected at least issue key + common workflow fields."
                )
            expected_columns = len(header)
            pending: list[tuple[int, list[str]]] = []

            def flush_pending():
                nonlocal pending
                rows = pending
                pending = []
                return rows

            for logical_row, values in enumerate(reader, 2):
                if not any(str(value or "").strip() for value in values):
                    continue
                values = list(values)
                if not pending and len(values) == expected_columns:
                    yield "", logical_row, header, values
                    continue

                if len(values) != expected_columns:
                    pending.append((logical_row, values))
                    # Deterministic recovery for a common broken TSV/Jira export
                    # shape: prefix columns + one unescaped multiline field +
                    # suffix columns. Middle fragments must each be exactly one
                    # field and the reconstructed width must equal the header.
                    if len(pending) >= 3:
                        first_row, first = pending[0]
                        last_row, last = pending[-1]
                        middle = pending[1:-1]
                        if (
                            len(first) >= 2
                            and len(last) >= 2
                            and all(len(fragment) == 1 for _, fragment in middle)
                            and len(first) + 1 + len(last) == expected_columns
                        ):
                            continuation = "\n".join(fragment[0] for _, fragment in middle)
                            repaired = SourceRowValues(
                                [*first, continuation, *last],
                                repair_info={
                                    "strategy": "PREFIX_MULTILINE_SUFFIX_EXACT_WIDTH",
                                    "source_logical_rows": [row for row, _ in pending],
                                    "original_widths": [len(row_values) for _, row_values in pending],
                                    "expected_columns": expected_columns,
                                },
                            )
                            pending = []
                            yield "", first_row, header, repaired
                    continue

                # A full-width row marks the end of an unrecoverable fragment
                # sequence. Preserve every fragment for quarantine/audit rather
                # than guessing missing column positions.
                for pending_row, pending_values in flush_pending():
                    yield "", pending_row, header, pending_values
                yield "", logical_row, header, values

            for pending_row, pending_values in flush_pending():
                yield "", pending_row, header, pending_values


class QABugImportPipeline:
    """Enterprise normalize/validate/JSONL/SQL pipeline.

    SQL and Qdrant operations are opt-in. Dry-run and JSONL export require no
    database services, making data mapping review possible before persistence.
    """

    def __init__(self, canonicalizer: QABugCanonicalizer | None = None) -> None:
        self.canonicalizer = canonicalizer or QABugCanonicalizer()

    @staticmethod
    def _reader(path: Path):
        suffix = path.suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            return QABugFileReader.iter_xlsx(path)
        if suffix in {".txt", ".tsv", ".csv"}:
            return QABugFileReader.iter_delimited(path)
        raise QABugImportError(f"Unsupported QA/BUG register format: {suffix}")

    def iter_documents(
        self,
        path: str | Path,
        *,
        record_type: str,
        max_records: int | None = None,
    ) -> Iterator[dict[str, Any]]:
        source = Path(path)
        if not source.is_file():
            raise QABugImportError(f"Input file not found: {source}")
        count = 0
        mapped_cache: dict[tuple[str, ...], list[tuple[str, str]]] = {}
        for sheet, row_number, headers, values in self._reader(source):
            header_key = tuple(str(x or "") for x in headers)
            mapped = mapped_cache.get(header_key)
            if mapped is None:
                mapped = self.canonicalizer.map_headers(headers)
                mapped_cache[header_key] = mapped
            document = self.canonicalizer.canonicalize(
                record_type=record_type,
                mapped_headers=mapped,
                values=values,
                source_file=str(source),
                source_sheet=sheet,
                source_row=row_number,
                source_format=source.suffix.lstrip(".").upper(),
            )
            if document is None:
                continue
            yield document
            count += 1
            if max_records is not None and count >= max_records:
                return

    def export_jsonl(
        self,
        path: str | Path,
        output_path: str | Path,
        *,
        record_type: str,
        max_records: int | None = None,
        progress: Callable[[int, str], None] | None = None,
    ) -> QABugImportReport:
        report = QABugImportReport(str(path), record_type.upper())
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        with out.open("w", encoding="utf-8", newline="\n") as handle:
            for document in self.iter_documents(path, record_type=record_type, max_records=max_records):
                report.rows_seen += 1
                report.rows_valid += 1
                if not bool(document.get("queryable", True)):
                    report.rows_quarantined += 1
                handle.write(json.dumps(document, ensure_ascii=False, separators=(",", ":")) + "\n")
                report.jsonl_written += 1
                if progress and report.rows_seen % 100 == 0:
                    progress(report.rows_seen, str(document.get("identity", {}).get("issue_key") or ""))
        return report

    def import_postgres(
        self,
        path: str | Path,
        *,
        record_type: str,
        profile,
        password: str,
        max_records: int | None = None,
        continue_on_error: bool = True,
        progress: Callable[[int, str], None] | None = None,
    ) -> QABugImportReport:
        from .repository import QABugRepository
        repo = QABugRepository(profile, password)
        report = QABugImportReport(str(path), record_type.upper())
        source = Path(path)
        digest = hashlib.sha256()
        with source.open("rb") as source_handle:
            for block in iter(lambda: source_handle.read(1024 * 1024), b""):
                digest.update(block)
        source_file_id, run_id = repo.begin_import(
            source_file=str(source), record_type=record_type.upper(), source_hash=digest.hexdigest(),
            file_size_bytes=source.stat().st_size, options={"max_records": max_records, "continue_on_error": continue_on_error},
        )
        try:
            for document in self.iter_documents(path, record_type=record_type, max_records=max_records):
                report.rows_seen += 1
                report.rows_valid += 1
                if not bool(document.get("queryable", True)):
                    report.rows_quarantined += 1
                key = str(document.get("identity", {}).get("issue_key") or document.get("identity", {}).get("record_key") or "")
                try:
                    _record_id, chunks = repo.upsert_document(document, source_file_id=source_file_id)
                    report.records_upserted += 1
                    report.chunks_upserted += int(chunks)
                except Exception as exc:
                    report.rows_failed += 1
                    report.errors.append(f"{key}: {exc}")
                    # Keep row-level failure provenance in SQL. Failure to write
                    # the audit row must never hide the original import error.
                    try:
                        repo.record_import_error(
                            run_id,
                            source_row_number=document.get("source", {}).get("row_number"),
                            issue_key=key,
                            error_code="UPSERT_FAILED",
                            error_message=str(exc),
                            raw_json={
                                "identity": document.get("identity", {}),
                                "source": document.get("source", {}),
                                "quality_flags": document.get("quality_flags", []),
                            },
                        )
                    except Exception as audit_exc:
                        report.errors.append(f"{key}: import-error audit persistence failed: {audit_exc}")
                    if not continue_on_error:
                        raise
                if progress and report.rows_seen % 25 == 0:
                    progress(report.rows_seen, key)
        finally:
            repo.finish_import(
                run_id, status="COMPLETED_WITH_ERRORS" if report.rows_failed else "COMPLETED",
                rows_seen=report.rows_seen, rows_valid=report.rows_valid, rows_upserted=report.records_upserted,
                rows_skipped=report.rows_skipped, rows_failed=report.rows_failed, chunks_upserted=report.chunks_upserted,
                error_summary="\n".join(report.errors[:20]),
            )
        return report

    def detect_record_type(self, path: str | Path, *, sample_rows: int = 80) -> tuple[str, float, dict[str, int]]:
        """Infer QA vs BUG from source semantics, never filename alone.

        Explicit import buttons remain authoritative. This helper is used by
        generic JSONL preview/export and provides reviewable scoring.
        """
        source = Path(path)
        qa_score = 0
        bug_score = 0
        seen = 0
        mapped_cache: dict[tuple[str, ...], list[tuple[str, str]]] = {}
        for _sheet, _row_number, headers, values in self._reader(source):
            header_key = tuple(str(x or "") for x in headers)
            mapped = mapped_cache.get(header_key)
            if mapped is None:
                mapped = self.canonicalizer.map_headers(headers)
                mapped_cache[header_key] = mapped
                canonical_headers = {canonical for _, canonical in mapped if canonical}
                qa_score += 3 * len(canonical_headers & {"answer", "final_answer", "rationale", "reference_spec", "procedure_expected"})
                bug_score += 3 * len(canonical_headers & {"root_cause", "countermeasure", "effect_result", "cause_module", "analysis_result"})
            values_map = self.canonicalizer._canonical_values(mapped, values)
            issue_type = str(values_map.get("issue_type") or "")
            if any(token in issue_type.lower() for token in ("仕様問題", "qa", "question", "问答", "問答")):
                qa_score += 6
            if any(token in issue_type.lower() for token in ("不具合", "bug", "defect", "障害")):
                bug_score += 6
            if any(values_map.get(k) for k in ("answer", "final_answer", "rationale", "reference_spec")):
                qa_score += 2
            if any(values_map.get(k) for k in ("root_cause", "analysis_result", "countermeasure", "effect_result")):
                bug_score += 2
            seen += 1
            if seen >= max(1, int(sample_rows)):
                break
        if seen == 0:
            raise QABugImportError(f"No data rows found in {source.name}")
        total = qa_score + bug_score
        if total == 0:
            return "BUG", 0.0, {"QA": qa_score, "BUG": bug_score}
        result = "QA" if qa_score > bug_score else "BUG"
        confidence = max(qa_score, bug_score) / total
        return result, round(confidence, 4), {"QA": qa_score, "BUG": bug_score}

    def field_mapping_report(self, path: str | Path) -> dict[str, Any]:
        source = Path(path)
        rows = self._reader(source)
        try:
            sheet, row_number, headers, _values = next(rows)
        except StopIteration:
            raise QABugImportError(f"No data rows found in {source.name}")
        mapped = self.canonicalizer.map_headers(headers)
        recognized = [{"source": raw, "canonical": canonical} for raw, canonical in mapped if canonical]
        unknown = [raw for raw, canonical in mapped if not canonical]
        return {
            "file": source.name,
            "sheet": sheet,
            "header_row_or_first_data_row": row_number - 1,
            "column_count": len(mapped),
            "recognized_count": len(recognized),
            "unknown_count": len(unknown),
            "recognized": recognized,
            "unknown": unknown,
            "unknown_policy": "PRESERVE_IN_RAW_JSON_AND_CUSTOM_FIELDS",
        }
